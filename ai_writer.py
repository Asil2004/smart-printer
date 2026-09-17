"""
ai_writer.py — AI yordamida hujjat yaratish moduli
====================================================
Gemini Flash modeli yordamida O'zbekiston standartlarida
rasmiy hujjatlar (ariza, dalolatnoma, mustaqil ish, qayta
o'zlashtirish) yaratadi va DOCX formatda saqlaydi.
"""

import os
import json
import base64
import logging
import tempfile
import urllib.request
from pathlib import Path
from datetime import datetime

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

logger = logging.getLogger("smart_printchi.ai_writer")

# ──────────────────────────────────────────────────────────
# Hujjat turlari va ularning shablonlari
# ──────────────────────────────────────────────────────────
DOC_TYPES = {
    "ariza": {
        "name":        "Ariza",
        "description": "Rasmiy ariza (ta'til, ko'chirish, imtihon va h.k.)",
        "base_price":  3000,
        "icon":        "fa-file-signature",
    },
    "dalolatnoma": {
        "name":        "Dalolatnoma",
        "description": "Rasmiy dalolatnoma / akt",
        "base_price":  4000,
        "icon":        "fa-file-contract",
    },
    "mustaqil": {
        "name":        "Mustaqil Ish",
        "description": "Akademik mustaqil ish / referat",
        "base_price":  5000,
        "icon":        "fa-graduation-cap",
    },
    "qayta": {
        "name":        "Qayta O'zlashtirish",
        "description": "Qayta o'zlashtirish uchun ariza / xat",
        "base_price":  3000,
        "icon":        "fa-rotate-right",
    },
    "xat": {
        "name":        "Rasmiy Xat",
        "description": "Muassasaga rasmiy xat / maktub",
        "base_price":  3000,
        "icon":        "fa-envelope-open-text",
    },
}

# So'z narxi (100 so'zdan oshgan har 100 so'z uchun)
PRICE_PER_100_WORDS = 500
FREE_WORDS_LIMIT    = 100


def _build_prompt(doc_type: str, recipient: str, content: str,
                  author_name: str, date: str, organization: str) -> str:
    """Hujjat turi va ma'lumotlarga qarab Gemini uchun prompt yaratadi."""

    today = date or datetime.now().strftime("%d.%m.%Y")

    base = f"""Siz O'zbekiston rasmiy hujjatlarini yozishda mutaxassis assistantsiz.
Quyidagi ma'lumotlar asosida "{DOC_TYPES.get(doc_type, {}).get('name', doc_type)}" turida
rasmiy hujjat yozing. Hujjat O'zbekiston davlat standartlariga (O'z DSt) mos bo'lishi shart.

MUHIM QOIDALAR:
- Faqat hujjat matnini yozing, ortiqcha tushuntirish YOZMA
- Rasmiy, professional uslubda yozing
- O'zbek tilida yozing
- Sana formati: {today}
- Imzo joyi qoldirilsin

"""

    if doc_type == "ariza":
        return base + f"""ARIZA FORMATI:
Kimga: {recipient}
Kim tomonidan: {author_name}
Muassasa: {organization or "[muassasa nomi]"}

Ariza mazmuni: {content}

Standart ariza tuzilmasi:
1. Yuqori o'ng burchak: Kimga (lavozim, F.I.O.)
2. Kim tomonidan (lavozim, F.I.O.)
3. Markazda: "ARIZA" sarlavhasi
4. Asosiy matn (iltimos, so'rov)
5. Sana va imzo

Iltimos, to'liq formatlangan ariza matnini yozing."""

    elif doc_type == "dalolatnoma":
        return base + f"""DALOLATNOMA FORMATI:
Tashkilot: {organization or "[tashkilot nomi]"}
Ishtirokchilar: {recipient}
Voqea/Holat: {content}
Sana: {today}

Standart dalolatnoma tuzilmasi:
1. "DALOLATNOMA" sarlavhasi (markazda)
2. Sana va joy
3. Komissiya a'zolari
4. Dalolatnoma predmeti
5. Aniqlangan holat/faktlar
6. Xulosa
7. Imzolar uchun joy

To'liq formatlangan dalolatnoma yozing."""

    elif doc_type == "mustaqil":
        return base + f"""MUSTAQIL ISH FORMATI:
Mavzu: {content}
Talaba: {author_name}
Fan/Kafedra: {recipient}
Muassasa: {organization or "[universitet nomi]"}
Sana: {today}

Mustaqil ish tuzilmasi:
1. Kirish (mavzuning dolzarbligi)
2. Asosiy qism (2-3 bo'lim)
3. Xulosa
4. Foydalanilgan adabiyotlar

Ilmiy uslubda, kamida 300 so'zlik mustaqil ish yozing."""

    elif doc_type == "qayta":
        return base + f"""QAYTA O'ZLASHTIRISH ARIZASI FORMATI:
Kimga: {recipient}
Talaba: {author_name}
Muassasa: {organization or "[universitet nomi]"}
Sabab: {content}

Standart qayta o'zlashtirish arizasi:
1. Yuqori o'ng burchak: Kimga
2. Kim tomonidan
3. "ARIZA" sarlavhasi
4. Qayta o'zlashtirish so'rovi va sababi
5. Tegishli fan va o'qituvchi
6. Sana va imzo

To'liq ariza matnini yozing."""

    elif doc_type == "xat":
        return base + f"""RASMIY XAT FORMATI:
Kimga: {recipient}
Kimdan: {author_name} ({organization or "[muassasa]"})
Mavzu: {content}
Sana: {today}

Rasmiy xat tuzilmasi:
1. Chiquvchi raqam va sana (yuqori chap)
2. Kimga (yuqori o'ng)
3. Xat mavzusi
4. Asosiy matn
5. Xulosa va iltimos
6. Imzo

To'liq rasmiy xat yozing."""

    return base + f"Mazmun: {content}\nKimga: {recipient}\nKimdan: {author_name}"


def calculate_price(doc_type: str, word_count: int, include_print: bool = True,
                    print_pages: int = 1, color: bool = False) -> dict:
    """Hujjat yaratish va chop etish narxini hisoblaydi."""
    base   = DOC_TYPES.get(doc_type, {}).get("base_price", 3000)
    extra  = max(0, word_count - FREE_WORDS_LIMIT) // 100 * PRICE_PER_100_WORDS
    ai_fee = base + extra

    print_fee = 0
    if include_print:
        price_per_page = 1500 if color else 500
        print_fee      = print_pages * price_per_page

    total = ai_fee + print_fee
    return {
        "ai_fee":       ai_fee,
        "print_fee":    print_fee,
        "total":        total,
        "word_count":   word_count,
        "breakdown": {
            "base":     base,
            "extra":    extra,
            "printing": print_fee,
        }
    }


# ──────────────────────────────────────────────────────────
# Gemini REST API orqali to'g'ridan-to'g'ri chaqirish
# (Katta SDK o'rnatish shart emas!)
# ──────────────────────────────────────────────────────────
def _call_gemini_rest(prompt: str, image_bytes: bytes = None, image_mime: str = "image/jpeg") -> str:
    """Gemini REST API orqali matn generatsiya qiladi."""
    api_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if not api_key:
        raise ValueError("GEMINI_API_KEY topilmadi! Iltimos, .env faylga GEMINI_API_KEY ni kiriting.")

    parts = []
    if image_bytes:
        b64 = base64.b64encode(image_bytes).decode("utf-8")
        parts.append({
            "inlineData": {
                "mimeType": image_mime,
                "data": b64
            }
        })
    parts.append({"text": prompt})

    payload = {
        "contents": [{"parts": parts}],
        "generationConfig": {
            "temperature": 0.3,
            "maxOutputTokens": 2048
        }
    }

    # Modellarni tartib bilan tekshirish
    models_to_try = [
        ("v1beta", "gemini-3.6-flash"),
        ("v1alpha", "gemini-3.8-flash"),
        ("v1beta", "gemini-2.0-flash"),
    ]

    last_err = None
    for ver, model_name in models_to_try:
        url = f"https://generativelanguage.googleapis.com/{ver}/models/{model_name}:generateContent?key={api_key}"
        try:
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"}
            )
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                candidates = data.get("candidates", [])
                if candidates and "content" in candidates[0]:
                    parts = candidates[0]["content"].get("parts", [])
                    if parts and "text" in parts[0]:
                        return parts[0]["text"].strip()
        except Exception as e:
            logger.warning(f"REST model {model_name} xatosi: {e}")
            last_err = e

    raise RuntimeError(f"Gemini API bilan bog'lanishda xatolik: {last_err}")


# ──────────────────────────────────────────────────────────
# Rasmdan matn olish (OCR via Gemini Vision)
# ──────────────────────────────────────────────────────────
def extract_text_from_image(image_bytes: bytes, mime_type: str = "image/jpeg") -> str:
    """Gemini Vision yordamida rasmdan matnni chiqaradi."""
    prompt = "Bu rasmdagi barcha matnni aniq o'qib yozing. Faqat matnni qaytaring, tushuntirish yozmang."
    return _call_gemini_rest(prompt=prompt, image_bytes=image_bytes, image_mime=mime_type)


# ──────────────────────────────────────────────────────────
# AI hujjat yaratish
# ──────────────────────────────────────────────────────────
def generate_document_text(doc_type: str, recipient: str, content: str,
                            author_name: str = "", date: str = "",
                            organization: str = "",
                            image_bytes: bytes = None,
                            image_mime: str = "image/jpeg") -> dict:
    """Gemini API yordamida hujjat matnini yaratadi."""

    # Agar rasm berilgan bo'lsa — avval OCR qilish
    if image_bytes:
        logger.info("Rasmdan matn olinmoqda (Gemini Vision)...")
        extracted = extract_text_from_image(image_bytes, image_mime)
        content   = extracted if extracted else content
        logger.info(f"Rasmdan olingan matn: {content[:100]}...")

    prompt = _build_prompt(
        doc_type=doc_type,
        recipient=recipient,
        content=content,
        author_name=author_name,
        date=date,
        organization=organization,
    )

    logger.info(f"Gemini API ga so'rov yuborilmoqda (doc_type={doc_type})...")
    generated_text = _call_gemini_rest(prompt=prompt)
    word_count     = len(generated_text.split())

    logger.info(f"Hujjat yaratildi: {word_count} so'z")
    return {
        "text":       generated_text,
        "word_count": word_count,
        "doc_type":   doc_type,
        "price":      calculate_price(doc_type, word_count),
    }


# ──────────────────────────────────────────────────────────
# DOCX yaratish (python-docx bo'lsa docx, bo'lmasa matn fayl)
# ──────────────────────────────────────────────────────────
def create_docx(text: str, doc_type: str, author_name: str = "",
                output_path: Path = None) -> Path:
    """Yaratilgan hujjat matnini DOCX formatida saqlaydi."""
    if output_path is None:
        tmp = tempfile.NamedTemporaryFile(
            suffix=".docx", delete=False,
            prefix=f"ai_{doc_type}_"
        )
        output_path = Path(tmp.name)
        tmp.close()

    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        section          = doc.sections[0]
        section.page_width   = Cm(21)
        section.page_height  = Cm(29.7)
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(1.5)

        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)

        paragraphs = text.split("\n")
        for para_text in paragraphs:
            stripped = para_text.strip()
            if not stripped:
                doc.add_paragraph("")
                continue

            para = doc.add_paragraph()
            run  = para.add_run(stripped)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

            if stripped.isupper() and len(stripped) < 50:
                run.bold = True
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif stripped.startswith(("1.", "2.", "3.", "4.", "5.", "6.")):
                run.bold = True
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            para.paragraph_format.space_after  = Pt(0)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.line_spacing = Pt(18)

        doc.save(str(output_path))
        logger.info(f"DOCX saqlandi: {output_path}")

    except ImportError:
        # python-docx yo'q bo'lsa matn sifatida saqlaymiz
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)
        logger.info(f"Matn fayl saqlandi: {output_path}")

    return output_path
